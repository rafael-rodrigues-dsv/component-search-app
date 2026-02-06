"""
Repository for TB_PLANILHA (spreadsheet)
"""
from src.infrastructure.repositories.access_repository import AccessRepository

class SpreadsheetRepository:
    def __init__(self):
        self._access = AccessRepository()

    def save_to_sheet(self, site_url: str, emails_str: str, telefones_str: str, endereco_str: str = None, distancia_km: float = None):
        """
        Salva empresa na TB_PLANILHA (para exportação Excel)

        Args:
            site_url: URL do site
            emails_str: Emails separados por ;
            telefones_str: Telefones separados por ;
            endereco_str: Endereço completo formatado
            distancia_km: Distância em km (se calculada)
        """
        conn = self._access._get_connection()
        cursor = conn.cursor()

        # Verificar se já existe
        cursor.execute("SELECT SITE FROM TB_PLANILHA WHERE SITE = ?", (site_url,))
        existing = cursor.fetchone()

        if existing:
            # Atualizar registro existente
            cursor.execute(
                "UPDATE TB_PLANILHA SET EMAIL = ?, TELEFONE = ?, ENDERECO = ?, DISTANCIA_KM = ? WHERE SITE = ?",
                (emails_str, telefones_str, endereco_str, distancia_km, site_url)
            )
        else:
            # Inserir novo registro
            cursor.execute(
                "INSERT INTO TB_PLANILHA (SITE, EMAIL, TELEFONE, ENDERECO, DISTANCIA_KM) VALUES (?, ?, ?, ?, ?)",
                (site_url, emails_str, telefones_str, endereco_str, distancia_km)
            )

        conn.commit()
        try:
            cursor.close()
        except Exception:
            pass

    def export_to_excel(self, excel_path: str):
        conn = self._access._get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT SITE, EMAIL, TELEFONE, ENDERECO, DISTANCIA_KM FROM TB_PLANILHA ORDER BY DISTANCIA_KM, SITE")
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Empresas"
        ws['A1'] = 'SITE'
        ws['B1'] = 'EMAIL'
        ws['C1'] = 'TELEFONE'
        ws['D1'] = 'ENDERECO'
        ws['E1'] = 'DISTANCIA_KM'
        row = 2
        for site, emails, telefones, endereco, distancia in cursor.fetchall():
            ws[f'A{row}'] = site
            ws[f'B{row}'] = emails or ''
            ws[f'C{row}'] = telefones or ''
            ws[f'D{row}'] = endereco or ''
            ws[f'E{row}'] = distancia or ''
            row += 1
        wb.save(excel_path)
        try:
            cursor.close()
        except Exception:
            pass
        return row - 2

    # --- NEW: pagination helpers ---
    def list_rows(self, limit: int = None, offset: int = 0):
        sql = "SELECT SITE AS site, EMAIL AS email, TELEFONE AS telefone, ENDERECO AS endereco, DISTANCIA_KM AS distancia_km FROM TB_PLANILHA ORDER BY DISTANCIA_KM, SITE"
        rows = self._access.execute_query(sql)
        if limit is None:
            return rows
        # Parse strictly
        limit = int(limit) if limit is not None else None
        offset = int(offset) if offset else 0
        return rows[offset: offset + limit] if rows else []

    def count(self) -> int:
        try:
            rows = self._access.execute_query("SELECT COUNT(*) as cnt FROM TB_PLANILHA")
            if isinstance(rows, list) and rows:
                return int(rows[0].get('cnt', 0) or 0)
            return 0
        except Exception:
            return 0
